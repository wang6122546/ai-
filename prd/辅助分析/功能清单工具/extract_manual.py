from docx import Document
import json

source = "/Users/wang/Desktop/重庆安全生产平台用户操作手册w.docx"
doc = Document(source)

paras = []
for idx, p in enumerate(doc.paragraphs):
    text = " ".join(p.text.split())
    if text:
        paras.append({"idx": idx, "style": p.style.name if p.style else "", "text": text})

tables = []
for ti, table in enumerate(doc.tables):
    rows = []
    for row in table.rows:
        rows.append([" ".join(cell.text.split()) for cell in row.cells])
    tables.append({"idx": ti, "rows": rows})

with open("manual_extract.json", "w", encoding="utf-8") as f:
    json.dump({"paragraphs": paras, "tables": tables}, f, ensure_ascii=False, indent=2)

print(json.dumps({"paragraphs": len(paras), "tables": len(tables)}, ensure_ascii=False))
